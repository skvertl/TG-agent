"""Document parsers and page-aware recursive character chunker."""
import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Tuple, Optional

from core.models import DocumentChunk
from core.observability.overlap import estimate_tokens


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


def _extract_docx_with_pages(content: bytes) -> List[Tuple[int, str]]:
    """
    Extracts text and page boundaries from DOCX content.
    1. If valid zip archive, inspects docProps/app.xml for metadata page count (<Pages>).
    2. Parses word/document.xml sequentially looking for:
       - Hard page breaks: <w:br w:type="page"/>
       - Word-rendered soft page breaks: <w:lastRenderedPageBreak/>
       - Paragraphs: <w:p>
       - Tables: <w:tbl>, rows <w:tr>, cells <w:tc>
       - Text: <w:t> and tabs <w:tab/>
    3. If XML page breaks yield >= 2 pages, uses the break-derived pages.
    4. If XML page breaks yield only 1 page, but metadata reports multiple pages (e.g. 86 pages),
       distributes paragraphs proportionally across the metadata pages.
    5. Fallback: if not a zip (e.g. mock in tests) or on parse error, uses python-docx.
    """
    if not zipfile.is_zipfile(io.BytesIO(content)):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text)
            return [(1, text)]
        except Exception:
            return [(1, "")]

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            expected_pages = 1
            if "docProps/app.xml" in z.namelist():
                try:
                    app_xml = z.read("docProps/app.xml")
                    root_app = ET.fromstring(app_xml)
                    for elem in root_app.iter():
                        if elem.tag.endswith("Pages") and elem.text and elem.text.strip().isdigit():
                            expected_pages = max(1, int(elem.text.strip()))
                except Exception:
                    pass

            if "word/document.xml" not in z.namelist():
                return [(1, "")]

            doc_xml = z.read("word/document.xml")
            root = ET.fromstring(doc_xml)
            W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

            pages: List[str] = []
            current_tokens: List[str] = []

            def flush_page():
                t = "".join(current_tokens).strip()
                if t:
                    pages.append(t)
                elif pages:
                    pages.append("")
                current_tokens.clear()

            body = root.find(f"{{{W_NS}}}body")
            if body is None:
                body = root

            for elem in body.iter():
                tag = elem.tag
                if tag == f"{{{W_NS}}}lastRenderedPageBreak" or (
                    tag == f"{{{W_NS}}}br" and elem.attrib.get(f"{{{W_NS}}}type") == "page"
                ):
                    flush_page()
                elif tag == f"{{{W_NS}}}t" and elem.text:
                    current_tokens.append(elem.text)
                elif tag == f"{{{W_NS}}}tab":
                    current_tokens.append("\t")
                elif tag == f"{{{W_NS}}}p" or tag == f"{{{W_NS}}}tr":
                    if current_tokens and not current_tokens[-1].endswith("\n"):
                        current_tokens.append("\n")

            flush_page()

            clean_pages = [(idx + 1, p) for idx, p in enumerate(pages) if p.strip()]
            if len(clean_pages) > 1:
                return clean_pages

            full_text = clean_pages[0][1] if clean_pages else ""
            if not full_text:
                return [(1, "")]

            if expected_pages > 1 and len(full_text) > expected_pages * 20:
                paragraphs = [p for p in full_text.split("\n") if p.strip()]
                total_chars = sum(len(p) for p in paragraphs)
                chars_per_page = max(1, total_chars // expected_pages)

                distributed_pages: List[Tuple[int, str]] = []
                curr_p_page: List[str] = []
                curr_len = 0
                current_page_num = 1

                for p in paragraphs:
                    curr_p_page.append(p)
                    curr_len += len(p)
                    if curr_len >= chars_per_page and current_page_num < expected_pages:
                        distributed_pages.append((current_page_num, "\n".join(curr_p_page)))
                        curr_p_page = []
                        curr_len = 0
                        current_page_num += 1

                if curr_p_page:
                    if distributed_pages and current_page_num <= expected_pages:
                        distributed_pages.append((current_page_num, "\n".join(curr_p_page)))
                    elif distributed_pages:
                        last_num, last_text = distributed_pages[-1]
                        distributed_pages[-1] = (last_num, last_text + "\n" + "\n".join(curr_p_page))
                    else:
                        distributed_pages.append((1, "\n".join(curr_p_page)))

                return distributed_pages

            return [(1, full_text)]
    except Exception:
        import docx
        doc = docx.Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs if p.text)
        return [(1, text)]


def extract_text_with_pages(content: bytes, filename: str) -> List[Tuple[int, str]]:
    """
    Extracts text from supported document formats (.txt, .md, .docx, .pdf),
    annotating each page with a 1-based page number.
    Returns list of (page_number, page_text).
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file format: '{ext}'. Supported formats are: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext in {".txt", ".md"}:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")
            except UnicodeDecodeError:
                text = content.decode("utf-8", errors="replace")
        return [(1, text)]

    elif ext == ".docx":
        return _extract_docx_with_pages(content)

    elif ext == ".pdf":
        try:
            import pypdf
        except ImportError:
            raise ImportError("pypdf is required to parse PDF documents. Please install pypdf.")
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages: List[Tuple[int, str]] = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            pages.append((idx + 1, page_text))
        return pages if pages else [(1, "")]

    raise ValueError(f"Unsupported file format: '{ext}'")


class RecursiveCharacterChunker:
    """
    Recursive character text splitter that breaks text into chunks based on a hierarchy
    of separators (paragraphs, newlines, sentences, words, characters), maintaining overlap.
    """

    def __init__(
        self,
        chunk_size: int = 600,
        overlap: int = 100,
        chunk_overlap: Optional[int] = None,
        separators: Optional[List[str]] = None,
    ) -> None:
        effective_overlap = chunk_overlap if chunk_overlap is not None else overlap
        if effective_overlap >= chunk_size:
            raise ValueError(f"overlap ({effective_overlap}) must be smaller than chunk_size ({chunk_size})")
        self.chunk_size = chunk_size
        self.overlap = effective_overlap
        self.separators = separators or ["\n\n", "\n", ". ", " ", ""]

    def split_text(self, text: str) -> List[str]:
        """Splits a single text block into chunks."""
        if not text or not text.strip():
            return []
        return self._split_text(text, self.separators)

    def _split_text(self, text: str, separators: List[str]) -> List[str]:
        final_chunks: List[str] = []
        separator = separators[-1]
        new_separators: List[str] = []
        for i, sep in enumerate(separators):
            if sep == "":
                separator = ""
                break
            if sep in text:
                separator = sep
                new_separators = separators[i + 1:]
                break

        splits = text.split(separator) if separator else list(text)
        good_splits: List[str] = []
        for s in splits:
            if not s:
                continue
            if len(s) < self.chunk_size:
                good_splits.append(s)
            else:
                if good_splits:
                    merged = self._merge_splits(good_splits, separator)
                    final_chunks.extend(merged)
                    good_splits = []
                if not new_separators:
                    final_chunks.append(s)
                else:
                    sub_chunks = self._split_text(s, new_separators)
                    final_chunks.extend(sub_chunks)
        if good_splits:
            merged = self._merge_splits(good_splits, separator)
            final_chunks.extend(merged)
        return [c.strip() for c in final_chunks if c.strip()]

    def _merge_splits(self, splits: List[str], separator: str) -> List[str]:
        docs: List[str] = []
        current_doc: List[str] = []
        total = 0
        for d in splits:
            _len = len(d)
            sep_len = len(separator) if current_doc else 0
            if total + _len + sep_len > self.chunk_size:
                if total > 0:
                    doc = separator.join(current_doc).strip()
                    if doc:
                        docs.append(doc)
                    while total > self.overlap or (total + _len + sep_len > self.chunk_size and total > 0):
                        total -= len(current_doc[0]) + (len(separator) if len(current_doc) > 1 else 0)
                        current_doc = current_doc[1:]
            current_doc.append(d)
            total += _len + (len(separator) if len(current_doc) > 1 else 0)
        if current_doc:
            doc = separator.join(current_doc).strip()
            if doc:
                docs.append(doc)
        return docs

    def chunk_document(
        self,
        user_id: str,
        pages: List[Tuple[int, str]],
        document_id: Optional[str] = None,
    ) -> List[DocumentChunk]:
        """
        Takes parsed pages [(page_num, text), ...] and returns a sequence of DocumentChunk models
        with 0-based chunk indices, page numbers, and estimated token counts.
        """
        doc_id = document_id or "doc-" + user_id
        chunks: List[DocumentChunk] = []
        chunk_index = 0
        for page_num, text in pages:
            page_chunks = self.split_text(text)
            for chunk_text in page_chunks:
                chunks.append(
                    DocumentChunk(
                        document_id=doc_id,
                        user_id=user_id,
                        content=chunk_text,
                        page_number=page_num,
                        chunk_index=chunk_index,
                        token_count=estimate_tokens(chunk_text),
                    )
                )
                chunk_index += 1
        return chunks

    def split_pages(
        self,
        pages: List[Tuple[int, str]],
        document_id: str,
        user_id: str,
    ) -> List[DocumentChunk]:
        """Splits page-annotated text, preserving page_number and continuity."""
        return self.chunk_document(user_id=user_id, pages=pages, document_id=document_id)
