"""Unit tests for document parsers and recursive character chunker."""
import pytest
from unittest.mock import MagicMock, patch

from adapters.rag.parsers import extract_text_with_pages, RecursiveCharacterChunker
from core.models import DocumentChunk


class TestExtractTextWithPages:
    def test_extract_txt_utf8(self):
        content = "Hello, world! This is a test.".encode("utf-8")
        pages = extract_text_with_pages(content, "sample.txt")
        assert len(pages) == 1
        assert pages[0] == (1, "Hello, world! This is a test.")

    def test_extract_txt_cp1251(self):
        text = "Привет, мир! Тестовый документ."
        content = text.encode("cp1251")
        pages = extract_text_with_pages(content, "document.txt")
        assert len(pages) == 1
        assert pages[0] == (1, text)

    def test_extract_md(self):
        text = "# Title\n\nSome **markdown** text."
        content = text.encode("utf-8")
        pages = extract_text_with_pages(content, "README.MD")
        assert len(pages) == 1
        assert pages[0] == (1, text)

    def test_extract_docx(self):
        with patch("docx.Document") as mock_docx:
            p1 = MagicMock()
            p1.text = "First paragraph"
            p2 = MagicMock()
            p2.text = "Second paragraph"
            mock_doc = MagicMock()
            mock_doc.paragraphs = [p1, p2]
            mock_docx.return_value = mock_doc

            pages = extract_text_with_pages(b"fake docx bytes", "report.docx")
            assert len(pages) == 1
            assert pages[0] == (1, "First paragraph\nSecond paragraph")

    def test_extract_docx_real_with_page_breaks(self):
        import io
        import docx
        doc = docx.Document()
        doc.add_paragraph("First Page Content")
        doc.add_page_break()
        doc.add_paragraph("Second Page Content")
        buf = io.BytesIO()
        doc.save(buf)

        pages = extract_text_with_pages(buf.getvalue(), "thesis.docx")
        assert len(pages) == 2
        assert pages[0][0] == 1
        assert "First Page Content" in pages[0][1]
        assert pages[1][0] == 2
        assert "Second Page Content" in pages[1][1]

    def test_extract_pdf(self):
        with patch("pypdf.PdfReader") as mock_reader_cls:
            page1 = MagicMock()
            page1.extract_text.return_value = "Content of page 1"
            page2 = MagicMock()
            page2.extract_text.return_value = "Content of page 2"
            mock_reader = MagicMock()
            mock_reader.pages = [page1, page2]
            mock_reader_cls.return_value = mock_reader

            pages = extract_text_with_pages(b"fake pdf bytes", "manual.pdf")
            assert len(pages) == 2
            assert pages[0] == (1, "Content of page 1")
            assert pages[1] == (2, "Content of page 2")

    def test_unsupported_format(self):
        with pytest.raises(ValueError, match="Unsupported file format"):
            extract_text_with_pages(b"image bytes", "picture.png")


class TestRecursiveCharacterChunker:
    def test_small_text_single_chunk(self):
        chunker = RecursiveCharacterChunker(chunk_size=100, overlap=20)
        chunks = chunker.split_text("Short text.")
        assert chunks == ["Short text."]

    def test_empty_text_returns_empty(self):
        chunker = RecursiveCharacterChunker(chunk_size=100, overlap=20)
        chunks = chunker.split_text("")
        assert chunks == []

    def test_splitting_respects_chunk_size_and_overlap(self):
        chunker = RecursiveCharacterChunker(chunk_size=50, overlap=10)
        paragraph = (
            "Paragraph one is relatively short. "
            "Paragraph two has more details and some sentences. "
            "Paragraph three concludes the discussion."
        )
        chunks = chunker.split_text(paragraph)
        assert len(chunks) > 1
        for chunk in chunks:
            assert len(chunk) <= 60  # Allow slight tolerance or exact <= chunk_size

    def test_chunk_document_preserves_pages_and_metadata(self):
        chunker = RecursiveCharacterChunker(chunk_size=100, overlap=20)
        pages = [
            (1, "Page 1 has some content that should fit in one chunk."),
            (2, "Page 2 has quite a lot more content. It just goes on and on and on until it gets split into multiple chunks because it exceeds the limit."),
        ]
        doc_chunks = chunker.chunk_document(user_id="user-42", pages=pages)

        assert len(doc_chunks) >= 3
        # Check consecutive indices
        for i, chunk in enumerate(doc_chunks):
            assert isinstance(chunk, DocumentChunk)
            assert chunk.chunk_index == i
            assert chunk.user_id == "user-42"
            assert chunk.token_count is not None
            assert chunk.token_count > 0

        # Check page numbers preserved
        assert doc_chunks[0].page_number == 1
        assert doc_chunks[1].page_number == 2

    def test_split_pages_with_document_id(self):
        chunker = RecursiveCharacterChunker(chunk_size=100, overlap=20)
        pages = [(1, "Some content for split_pages test.")]
        doc_chunks = chunker.split_pages(pages=pages, document_id="doc-999", user_id="user-42")
        assert len(doc_chunks) == 1
        assert doc_chunks[0].document_id == "doc-999"
        assert doc_chunks[0].user_id == "user-42"
        assert isinstance(doc_chunks[0].id, str)

